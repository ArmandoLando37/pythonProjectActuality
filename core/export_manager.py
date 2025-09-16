import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
import os

class ExportManager:
    def export_to_excel(self, articles, filename="news_analysis.xlsx"):
        """Exporte les articles vers un fichier Excel avec graphiques"""
        try:
            df = self.create_dataframe(articles)
            if not df.empty:
                with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                    # Feuille Articles
                    df.to_excel(writer, sheet_name='Articles', index=False)
                    
                    # Feuille Statistiques
                    workbook = writer.book
                    self.create_statistics_sheet(workbook, df)
                    
                return True, f"Export réussi : {filename}"
            return False, "Aucune donnée à exporter"
        except Exception as e:
            return False, f"Erreur lors de l'export : {str(e)}"
    
    def create_dataframe(self, articles):
        """Crée un DataFrame à partir des articles"""
        import pandas as pd
        from datetime import datetime
        
        data = []
        for article in articles:
            data.append({
                'Titre': article.get('title', ''),
                'Source': article.get('source', ''),
                'Date': article.get('published_at', ''),
                'Auteur': article.get('author', ''),
                'Description': article.get('description', ''),
                'URL': article.get('url', '')
            })
        
        return pd.DataFrame(data)
    
    def create_statistics_sheet(self, workbook, df):
        """Crée la feuille de statistiques avec graphiques"""
        worksheet = workbook.create_sheet('Statistiques')
        
        # Statistiques par source
        source_stats = df['Source'].value_counts().reset_index()
        source_stats.columns = ['Source', 'Nombre d\'articles']
        
        # Écrire les données
        worksheet['A1'] = 'Statistiques par Source'
        for col_num, value in enumerate(source_stats.columns, 1):
            worksheet.cell(row=3, column=col_num, value=value)
        
        for row_num, (source, count) in enumerate(source_stats.values, 4):
            worksheet.cell(row=row_num, column=1, value=source)
            worksheet.cell(row=row_num, column=2, value=count)
        
        # Créer un graphique
        chart = BarChart()
        chart.type = "col"
        chart.title = "Articles par Source"
        chart.x_axis.title = "Sources"
        chart.y_axis.title = "Nombre d'articles"
        
        data = Reference(worksheet, min_col=2, min_row=3, max_row=3 + len(source_stats))
        categories = Reference(worksheet, min_col=1, min_row=4, max_row=3 + len(source_stats))
        
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        
        worksheet.add_chart(chart, "D3")